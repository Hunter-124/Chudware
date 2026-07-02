"""Template registry + the three smoke templates (contract §8).

A **template** is a named recipe that maps outline blocks to a file via one engine, and — the
load-bearing bit — declares which slot names / workbook columns are *load-bearing* per block
type (contract §4.3), so the validator can reject a literal where a provenance ref is required.
Templates are looked up by name through a generic registry, so real engineering templates
(T-F15) slot in with no contract change.

``deliverables/1`` ships exactly one smoke template per format to prove the seam end-to-end:
``smoke_deck`` (pptx), ``smoke_report`` (docx/pdf), ``smoke_workbook`` (xlsx, live formula).

Every engine import is lazy (inside the builder) so importing this module — e.g. for
``list_artifacts`` — never requires python-pptx/docx/openpyxl/reportlab to be installed.
"""
from __future__ import annotations

from pathlib import Path
from typing import Callable

from .errors import DeliverableError

# The engineering result slots the smoke deck/report treat as load-bearing (§4.3): a literal
# in any of these fails validation with provenance_violation.
_RESULT_SLOTS = {"max_stress", "min_fos", "mass"}


class Template:
    """A named recipe. ``build(resolved, fmt, out_path)`` writes the file with one engine."""

    def __init__(self, name: str, kind: str,
                 build: Callable[[dict, str, Path], None],
                 lb_slots: dict[str, set] | None = None,
                 lb_cols: dict[str, set] | None = None) -> None:
        self.name = name
        self.kind = kind
        self._build = build
        self._lb_slots = lb_slots or {}
        self._lb_cols = lb_cols or {}

    def load_bearing_slots(self, block_type: str) -> set:
        return self._lb_slots.get(block_type, set())

    def load_bearing_columns(self, block_type: str) -> set:
        return self._lb_cols.get(block_type, set())

    def build(self, resolved: dict, fmt: str, out_path: Path) -> None:
        self._build(resolved, fmt, out_path)


_REGISTRY: dict[str, Template] = {}


def register_template(t: Template) -> None:
    _REGISTRY[t.name] = t


def get_template(name: str, kind: str) -> Template:
    t = _REGISTRY.get(name)
    if t is None:
        raise DeliverableError("invalid_spec", f"unknown template {name!r}")
    if t.kind != kind:
        raise DeliverableError(
            "invalid_spec", f"template {name!r} builds {t.kind!r}, not {kind!r}")
    return t


def _slot_lines(block: dict) -> list[str]:
    """Human ``label: value`` lines for a block's resolved slots (numbers came from tool results)."""
    out = []
    for name, rv in block.get("slots", {}).items():
        label = name.replace("_", " ").strip().capitalize()
        out.append(f"{label}: {rv['text']}")
    return out


# --------------------------------------------------------------------------- smoke_deck (pptx)

def _build_deck(resolved: dict, fmt: str, out_path: Path) -> None:
    from pptx import Presentation  # noqa: PLC0415

    prs = Presentation()
    blocks = resolved["blocks"] or [{"type": "title", "heading": resolved.get("title", ""),
                                     "body": [], "slots": {}, "rows": []}]
    for blk in blocks:
        if blk["type"] == "title":
            slide = prs.slides.add_slide(prs.slide_layouts[0])
            slide.shapes.title.text = blk.get("heading") or resolved.get("title", "")
            if len(slide.placeholders) > 1:
                slide.placeholders[1].text = "\n".join(blk.get("body", []))
        else:
            slide = prs.slides.add_slide(prs.slide_layouts[1])
            slide.shapes.title.text = blk.get("heading", "")
            body_lines = list(blk.get("body", [])) + _slot_lines(blk)
            if len(slide.placeholders) > 1:
                slide.placeholders[1].text = "\n".join(body_lines)
    prs.save(str(out_path))


# --------------------------------------------------------------------------- smoke_report (docx/pdf)

def _build_report(resolved: dict, fmt: str, out_path: Path) -> None:
    if fmt == "pdf":
        _build_report_pdf(resolved, out_path)
    else:
        _build_report_docx(resolved, out_path)


def _build_report_docx(resolved: dict, out_path: Path) -> None:
    from docx import Document  # noqa: PLC0415

    doc = Document()
    for blk in resolved["blocks"]:
        if blk["type"] == "title":
            doc.add_heading(blk.get("heading") or resolved.get("title", ""), level=0)
        else:
            doc.add_heading(blk.get("heading", ""), level=1)
        for line in blk.get("body", []):
            doc.add_paragraph(line)
        for line in _slot_lines(blk):
            doc.add_paragraph(line)
    doc.save(str(out_path))


def _build_report_pdf(resolved: dict, out_path: Path) -> None:
    # reportlab (BSD, in-process) — no subprocess, no LibreOffice needed for PDF reports.
    from reportlab.lib.pagesizes import letter  # noqa: PLC0415
    from reportlab.lib.styles import getSampleStyleSheet  # noqa: PLC0415
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer  # noqa: PLC0415

    styles = getSampleStyleSheet()
    flow = []
    for blk in resolved["blocks"]:
        style = styles["Title"] if blk["type"] == "title" else styles["Heading2"]
        flow.append(Paragraph(blk.get("heading") or resolved.get("title", ""), style))
        for line in blk.get("body", []):
            flow.append(Paragraph(line, styles["BodyText"]))
        for line in _slot_lines(blk):
            flow.append(Paragraph(line, styles["BodyText"]))
        flow.append(Spacer(1, 12))
    SimpleDocTemplate(str(out_path), pagesize=letter).build(flow)


# --------------------------------------------------------------------------- smoke_workbook (xlsx)

def _build_workbook(resolved: dict, fmt: str, out_path: Path) -> None:
    from openpyxl import Workbook  # noqa: PLC0415

    wb = Workbook()
    ws = wb.active
    first = True
    for blk in resolved["blocks"]:
        if blk["type"] != "sheet":
            continue
        title = (blk.get("heading") or "Sheet")[:31]
        if first:
            ws.title = title
            first = False
        else:
            ws = wb.create_sheet(title=title)
        for r, row in enumerate(blk.get("rows", []), start=1):
            for c, cell in enumerate(row, start=1):
                target = ws.cell(row=r, column=c)
                if "formula" in cell:
                    target.value = cell["formula"]  # a real live formula (recomputes on open)
                elif isinstance(cell.get("raw"), (int, float, bool)):
                    target.value = cell["raw"]        # keep numeric type
                else:
                    target.value = cell.get("text", "")
    wb.save(str(out_path))


def _register_smoke_templates() -> None:
    register_template(Template(
        "smoke_deck", "presentation", _build_deck,
        lb_slots={"section": set(_RESULT_SLOTS)}))
    register_template(Template(
        "smoke_report", "report", _build_report,
        lb_slots={"section": set(_RESULT_SLOTS)}))
    register_template(Template(
        "smoke_workbook", "workbook", _build_workbook,
        lb_cols={"sheet": {1}}))  # column index 1 (the mass column) is load-bearing


_register_smoke_templates()
