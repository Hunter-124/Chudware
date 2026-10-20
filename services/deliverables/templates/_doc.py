"""A tiny format-agnostic document model shared by the report templates.

The three engineering report templates (``calc_package``, ``verification_report``,
``assembly_manual``) all target ``make_report``'s two formats — ``.docx`` (python-docx) and
``.pdf`` (reportlab, in-process). Rather than write each template twice, a template builds one
:class:`Doc` of ordered ops (headings, paragraphs, bullet lists, tables, key/value blocks,
images) and this module renders that Doc to either backend. Both engines are permissive and
in-process (contract §7); imports are lazy so importing a template never forces them.

Load-bearing numbers are already resolved to strings by the time they reach a Doc — a template
passes ``block['slots'][name]['text']`` in, never a value it typed.
"""
from __future__ import annotations

from pathlib import Path
from xml.sax.saxutils import escape


class Doc:
    """An ordered list of render ops. Templates call the ``add_*`` methods; a backend renders it."""

    def __init__(self, title: str = "") -> None:
        self.ops: list[tuple] = []
        if title:
            self.heading(title, level=0)

    def heading(self, text: str, level: int = 1) -> "Doc":
        self.ops.append(("heading", str(text), int(level)))
        return self

    def para(self, text: str) -> "Doc":
        if text:
            self.ops.append(("para", str(text)))
        return self

    def bullets(self, items) -> "Doc":
        items = [str(i) for i in (items or []) if str(i).strip()]
        if items:
            self.ops.append(("bullets", items))
        return self

    def kv(self, pairs) -> "Doc":
        """A key/value block: list of (label, value) — values are already-resolved strings."""
        pairs = [(str(k), str(v)) for k, v in (pairs or [])]
        if pairs:
            self.ops.append(("kv", pairs))
        return self

    def table(self, rows, header: bool = True) -> "Doc":
        """A grid of stringified cells; ``rows[0]`` is a header row when ``header`` is True."""
        grid = [[("" if c is None else str(c)) for c in row] for row in (rows or [])]
        if grid:
            self.ops.append(("table", grid, bool(header)))
        return self

    def image(self, path: Path | None, caption: str = "") -> "Doc":
        """Embed an image if the path resolved to a real file, else just its caption."""
        self.ops.append(("image", path, str(caption or "")))
        return self

    def spacer(self) -> "Doc":
        self.ops.append(("spacer",))
        return self


def render(doc: Doc, fmt: str, out_path: Path) -> None:
    if fmt == "pdf":
        _render_pdf(doc, out_path)
    else:
        _render_docx(doc, out_path)


# --------------------------------------------------------------------------- docx backend

def _render_docx(doc: Doc, out_path: Path) -> None:
    from docx import Document  # noqa: PLC0415
    from docx.shared import Inches  # noqa: PLC0415

    d = Document()
    for op in doc.ops:
        kind = op[0]
        if kind == "heading":
            d.add_heading(op[1], level=min(op[2], 9))
        elif kind == "para":
            d.add_paragraph(op[1])
        elif kind == "bullets":
            for item in op[1]:
                d.add_paragraph(item, style="List Bullet")
        elif kind == "kv":
            for label, value in op[1]:
                p = d.add_paragraph()
                p.add_run(f"{label}: ").bold = True
                p.add_run(value)
        elif kind == "table":
            grid, header = op[1], op[2]
            ncols = max((len(r) for r in grid), default=1)
            tbl = d.add_table(rows=0, cols=ncols)
            try:
                tbl.style = "Light Grid Accent 1"
            except Exception:  # noqa: BLE001 - style set is theme-dependent; plain is fine
                pass
            for ri, row in enumerate(grid):
                cells = tbl.add_row().cells
                for ci in range(ncols):
                    text = row[ci] if ci < len(row) else ""
                    cells[ci].text = text
                    if header and ri == 0:
                        for para in cells[ci].paragraphs:
                            for run in para.runs:
                                run.bold = True
        elif kind == "image":
            path, caption = op[1], op[2]
            if path is not None:
                try:
                    d.add_picture(str(path), width=Inches(5.5))
                except Exception:  # noqa: BLE001 - a broken image never breaks the report
                    d.add_paragraph(f"[render unavailable: {Path(path).name}]")
            if caption:
                cp = d.add_paragraph(caption)
                cp.italic = True
        elif kind == "spacer":
            d.add_paragraph("")
    d.save(str(out_path))


# --------------------------------------------------------------------------- pdf backend

def _render_pdf(doc: Doc, out_path: Path) -> None:
    from reportlab.lib import colors  # noqa: PLC0415
    from reportlab.lib.pagesizes import letter  # noqa: PLC0415
    from reportlab.lib.styles import getSampleStyleSheet  # noqa: PLC0415
    from reportlab.lib.units import inch  # noqa: PLC0415
    from reportlab.platypus import (  # noqa: PLC0415
        Image, ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle,
    )

    styles = getSampleStyleSheet()
    h = {0: styles["Title"], 1: styles["Heading1"], 2: styles["Heading2"]}
    flow = []
    for op in doc.ops:
        kind = op[0]
        if kind == "heading":
            flow.append(Paragraph(escape(op[1]), h.get(op[2], styles["Heading3"])))
        elif kind == "para":
            flow.append(Paragraph(escape(op[1]), styles["BodyText"]))
        elif kind == "bullets":
            flow.append(ListFlowable(
                [ListItem(Paragraph(escape(i), styles["BodyText"])) for i in op[1]], bulletType="bullet"))
        elif kind == "kv":
            for label, value in op[1]:
                flow.append(Paragraph(f"<b>{escape(label)}:</b> {escape(value)}", styles["BodyText"]))
        elif kind == "table":
            grid, header = op[1], op[2]
            data = [[Paragraph(escape(c), styles["BodyText"]) for c in row] for row in grid]
            tbl = Table(data, hAlign="LEFT")
            ts = [("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                  ("VALIGN", (0, 0), (-1, -1), "TOP")]
            if header:
                ts.append(("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey))
                ts.append(("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"))
            tbl.setStyle(TableStyle(ts))
            flow.append(tbl)
        elif kind == "image":
            path, caption = op[1], op[2]
            if path is not None:
                try:
                    flow.append(Image(str(path), width=5.0 * inch, height=3.5 * inch,
                                      kind="proportional"))
                except Exception:  # noqa: BLE001
                    flow.append(Paragraph(escape(f"[render unavailable: {Path(path).name}]"),
                                          styles["BodyText"]))
            if caption:
                flow.append(Paragraph(f"<i>{escape(caption)}</i>", styles["BodyText"]))
        elif kind == "spacer":
            flow.append(Spacer(1, 12))
    SimpleDocTemplate(str(out_path), pagesize=letter).build(flow)
