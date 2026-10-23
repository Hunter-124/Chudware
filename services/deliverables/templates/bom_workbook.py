"""``bom_workbook`` — the BOM / cost workbook (xlsx), an engineering template.

A bill of materials with **live cost formulas**: quantity × unit-cost extended cost and a
per-level rollup that recompute when the sheet is opened (openpyxl real formulas, contract §5.3).
One sheet per assembly level. Fixed column order so load-bearing columns are unambiguous:

    | Item (label) | Qty | Unit Cost | Ext Cost (=Qty*Unit) | Source |
       col 0          1        2            3 (formula)          4

- ``Qty`` (col 1) and ``Unit Cost`` (col 2) are **load-bearing** — a bare numeral there is a
  ``provenance_violation``; they resolve from a BOM / sourced-part tool result (contract §4).
- ``Ext Cost`` (col 3) is a live formula string like ``=B2*C2``; a rollup row uses ``=SUM(...)``.
- ``Source`` (col 4) is free prose — a sourced-part link when sourcing data is present, blank without
  (graceful: the workbook still builds).

Registered ``kind='workbook'`` → driven by the existing ``make_workbook`` tool, no contract
change.
"""
from __future__ import annotations

from pathlib import Path

# Column order is part of the template's contract with the model (see module docstring).
_QTY_COL = 1
_UNIT_COST_COL = 2


def _build(resolved: dict, fmt: str, out_path: Path) -> None:
    from openpyxl import Workbook  # noqa: PLC0415
    from openpyxl.styles import Font  # noqa: PLC0415

    wb = Workbook()
    ws = wb.active
    first = True
    for blk in resolved["blocks"]:
        if blk.get("type") != "sheet":
            continue
        title = (blk.get("heading") or "Level")[:31]
        if first:
            ws.title = title
            first = False
        else:
            ws = wb.create_sheet(title=title)

        for r, row in enumerate(blk.get("rows", []), start=1):
            for c, cell in enumerate(row, start=1):
                target = ws.cell(row=r, column=c)
                if "formula" in cell:
                    target.value = cell["formula"]            # live formula (recomputes on open)
                elif isinstance(cell.get("raw"), (int, float, bool)):
                    target.value = cell["raw"]                # keep numeric type
                else:
                    target.value = cell.get("text", "")
                if r == 1:
                    target.font = Font(bold=True)             # header row
        if ws.max_row >= 1:
            ws.freeze_panes = "A2"
        for col_cells in ws.columns:
            width = max((len(str(c.value)) for c in col_cells if c.value is not None), default=8)
            ws.column_dimensions[col_cells[0].column_letter].width = min(max(width + 2, 10), 48)

    wb.save(str(out_path))


from . import Template, register_template  # noqa: E402 - after builder to mirror sibling modules

register_template(Template(
    "bom_workbook", "workbook", _build,
    lb_cols={"sheet": {_QTY_COL, _UNIT_COST_COL}}))
