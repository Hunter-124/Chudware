"""The three ``deliverables/1`` smoke templates (contract §8).

One per format, proving the seam end-to-end: ``smoke_deck`` (pptx), ``smoke_report``
(docx/pdf), ``smoke_workbook`` (xlsx, with a live formula). Moved verbatim from the original
flat ``templates.py`` when the package was split for the engineering templates.
"""
from __future__ import annotations

from pathlib import Path
from xml.sax.saxutils import escape

from . import Template, register_template, slot_lines

# The engineering result slots the smoke deck/report treat as load-bearing (§4.3): a literal
# in any of these fails validation with provenance_violation.
_RESULT_SLOTS = {"max_stress", "min_fos", "mass"}


# --------------------------------------------------------------------------- smoke_deck (pptx)

def _build_deck(resolved: dict, fmt: str, out_path: Path) -> None:
    from pptx import Presentation  # noqa: PLC0415

    prs = Presentation()
    blocks = resolved["blocks"] or [{"type": "title", "heading": resolved.get("title", ""),
                                     "body": [], "slots": {}, "rows": []}]
    for blk in blocks:
        if blk["type"] == "title":
            slide = prs.slides.add_slide(prs.slide_layouts[0])
            slide.shapes.title.text = blk.get("heading") or resolved.get("title", "")
            if len(slide.placeholders) > 1:
                slide.placeholders[1].text = "\n".join(blk.get("body", []))
        else:
            slide = prs.slides.add_slide(prs.slide_layouts[1])
            slide.shapes.title.text = blk.get("heading", "")
            body_lines = list(blk.get("body", [])) + slot_lines(blk)
            if len(slide.placeholders) > 1:
                slide.placeholders[1].text = "\n".join(body_lines)
    prs.save(str(out_path))


# --------------------------------------------------------------------------- smoke_report (docx/pdf)

def _build_report(resolved: dict, fmt: str, out_path: Path) -> None:
    if fmt == "pdf":
        _build_report_pdf(resolved, out_path)
    else:
        _build_report_docx(resolved, out_path)


def _build_report_docx(resolved: dict, out_path: Path) -> None:
    from docx import Document  # noqa: PLC0415

    doc = Document()
    for blk in resolved["blocks"]:
        if blk["type"] == "title":
            doc.add_heading(blk.get("heading") or resolved.get("title", ""), level=0)
        else:
            doc.add_heading(blk.get("heading", ""), level=1)
        for line in blk.get("body", []):
            doc.add_paragraph(line)
        for line in slot_lines(blk):
            doc.add_paragraph(line)
    doc.save(str(out_path))


def _build_report_pdf(resolved: dict, out_path: Path) -> None:
    # reportlab (BSD, in-process) — no subprocess, no LibreOffice needed for PDF reports.
    from reportlab.lib.pagesizes import letter  # noqa: PLC0415
    from reportlab.lib.styles import getSampleStyleSheet  # noqa: PLC0415
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer  # noqa: PLC0415

    styles = getSampleStyleSheet()
    flow = []
    for blk in resolved["blocks"]:
        style = styles["Title"] if blk["type"] == "title" else styles["Heading2"]
        flow.append(Paragraph(escape(blk.get("heading") or resolved.get("title", "")), style))
        for line in blk.get("body", []):
            flow.append(Paragraph(escape(line), styles["BodyText"]))
        for line in slot_lines(blk):
            flow.append(Paragraph(escape(line), styles["BodyText"]))
        flow.append(Spacer(1, 12))
    SimpleDocTemplate(str(out_path), pagesize=letter).build(flow)


# --------------------------------------------------------------------------- smoke_workbook (xlsx)

def _build_workbook(resolved: dict, fmt: str, out_path: Path) -> None:
    from openpyxl import Workbook  # noqa: PLC0415

    wb = Workbook()
    ws = wb.active
    first = True
    for blk in resolved["blocks"]:
        if blk["type"] != "sheet":
            continue
        title = (blk.get("heading") or "Sheet")[:31]
        if first:
            ws.title = title
            first = False
        else:
            ws = wb.create_sheet(title=title)
        for r, row in enumerate(blk.get("rows", []), start=1):
            for c, cell in enumerate(row, start=1):
                target = ws.cell(row=r, column=c)
                if "formula" in cell:
                    target.value = cell["formula"]  # a real live formula (recomputes on open)
                elif isinstance(cell.get("raw"), (int, float, bool)):
                    target.value = cell["raw"]        # keep numeric type
                else:
                    target.value = cell.get("text", "")
    wb.save(str(out_path))


register_template(Template(
    "smoke_deck", "presentation", _build_deck,
    lb_slots={"section": set(_RESULT_SLOTS)}))
register_template(Template(
    "smoke_report", "report", _build_report,
    lb_slots={"section": set(_RESULT_SLOTS)}))
register_template(Template(
    "smoke_workbook", "workbook", _build_workbook,
    lb_cols={"sheet": {1}}))  # column index 1 (the mass column) is load-bearing
