"""Tests for the deliverables tool family (contract ``deliverables/1``).

Covers:
  (a) a 4B-model-shaped spec (contract §5.1) producing a valid file for EACH format;
  (b) a bare numeral in a load-bearing slot is a validation ERROR (``provenance_violation``);
  (c) ``slot_unresolved`` cases (missing result / missing key / ambiguous tool);
  (d) ``render_document`` degrades cleanly without LibreOffice on PATH, and works with a
      *mocked* subprocess converter — so CI passes with neither pandoc nor LibreOffice installed;
  (e) artifact-registry round-trip + ``list_artifacts`` run_id filter.

The in-process engines (python-pptx/docx/openpyxl/matplotlib/reportlab/pypdfium2) are real; the
LibreOffice subprocess is mocked. Skipped cleanly when the in-process engines aren't installed.

Run:  python3 -m pytest services/deliverables/test_deliverables.py -q
"""
from __future__ import annotations

from pathlib import Path

import pytest


# The in-process engines are required deps; skip the whole module cleanly if a runtime lacks them.
pytest.importorskip("pptx")
pytest.importorskip("docx")
pytest.importorskip("openpyxl")
pytest.importorskip("matplotlib")
pytest.importorskip("reportlab")
pytest.importorskip("pypdfium2")

import services.deliverables as d  # noqa: E402
from services.deliverables import engines, project  # noqa: E402


@pytest.fixture(autouse=True)
def _project(tmp_path, monkeypatch):
    """Point BOTH roots at a fresh tmp dir per test (project.py reads the env live).

    Artifacts default to a hidden per-project store; pinning CHUDWARE_ARTIFACTS_DIR
    keeps the suite hermetic instead of writing into the real app data root.
    """
    monkeypatch.setenv("CHUDWARE_PROJECT_DIR", str(tmp_path))
    monkeypatch.setenv("CHUDWARE_ARTIFACTS_DIR", str(tmp_path / "artifacts"))
    yield


# --------------------------------------------------------------------------- fixtures

_ATTACHED = {
    "fea1": {"tool": "run_fea_static", "ts": "2026-07-02T18:20:41Z",
             "payload": {"max_von_mises_mpa": 182.4, "min_fos": 1.51, "max_disp_mm": 0.31}},
    "mass1": {"tool": "eval_mass", "ts": "2026-07-02T18:19:03Z",
              "payload": {"mass_g": 214.8, "material": "AL_6061_T6"}},
}


def _result_slots():
    return {
        "max_stress": {"from": {"tool": "run_fea_static", "key": "max_von_mises_mpa",
                                "result": "fea1"}, "unit": "MPa", "round": 1},
        "min_fos": {"from": {"tool": "run_fea_static", "key": "min_fos", "result": "fea1"},
                    "round": 2},
        "mass": {"from": {"tool": "eval_mass", "key": "mass_g", "result": "mass1"},
                 "unit": "g", "round": 1},
    }


def _report_spec(fmt="docx"):
    # The contract §5.1 spec: the model wrote structure + prose; code fills the numbers.
    return {
        "kind": "report",
        "title": "Bracket — Static Strength Report",
        "template": "smoke_report",
        "format": fmt,
        "run_id": "run_2026-07-02_bracket",
        "outline": [
            {"type": "title", "heading": "Bracket — Static Strength Report",
             "body": ["Prepared for the day-7 design review."]},
            {"type": "section", "heading": "Result summary",
             "body": ["Under the worst-case 1200 N tip load the bracket keeps a positive margin."],
             "slots": _result_slots()},
        ],
        "attached_results": _ATTACHED,
    }


def _deck_spec():
    return {
        "kind": "presentation", "title": "Bracket — Design Review", "template": "smoke_deck",
        "run_id": "run_2026-07-02_bracket",
        "outline": [
            {"type": "title", "heading": "Bracket — Design Review", "body": ["Day-7 review."]},
            {"type": "section", "heading": "Static FEA", "body": ["1200 N tip load."],
             "slots": _result_slots()},
        ],
        "attached_results": _ATTACHED,
    }


def _workbook_spec():
    return {
        "kind": "workbook", "title": "Bracket BOM", "template": "smoke_workbook",
        "run_id": "run_2026-07-02_bracket",
        "outline": [
            {"type": "sheet", "heading": "Summary", "rows": [
                ["Item", "Mass (g)", "Qty", "Total (g)"],
                ["Bracket",
                 {"from": {"tool": "eval_mass", "key": "mass_g", "result": "mass1"}, "round": 1},
                 2, "=B2*C2"],
            ]},
        ],
        "attached_results": _ATTACHED,
    }


def _abs(rec):
    # resolve_project_path is the documented inverse of the to_project_rel that
    # produced rec["path"] — exercising it here also covers the round trip.
    return project.resolve_project_path(rec["path"])


# ---------------------------------------------------------- (a) 4B-shaped spec, every format

_RECORD_KEYS = {"ok", "artifact_id", "kind", "path", "created_ts", "run_id", "source_refs",
                "spec_ref"}


def test_make_report_docx_full_roundtrip():
    rec = d.make_report(_report_spec("docx"))
    assert rec["ok"] is True
    assert _RECORD_KEYS <= set(rec)                      # frozen §3 record shape
    assert rec["kind"] == "report"
    assert rec["path"].endswith(".docx")
    assert _abs(rec).exists() and _abs(rec).stat().st_size > 0
    # every load-bearing number resolved from attached_results — none typed by the model.
    assert len(rec["source_refs"]) == 3
    assert {r["key"] for r in rec["source_refs"]} == {"max_von_mises_mpa", "min_fos", "mass_g"}
    # spec persisted verbatim for reproducibility.
    spec_abs = project.resolve_project_path(rec["spec_ref"])
    assert spec_abs.exists()
    # the docx really contains the resolved numbers.
    from docx import Document
    text = "\n".join(p.text for p in Document(str(_abs(rec))).paragraphs)
    assert "182.4" in text and "1.51" in text and "214.8" in text


def test_make_report_pdf():
    rec = d.make_report(_report_spec("pdf"))
    assert rec["ok"] is True and rec["path"].endswith(".pdf")
    assert _abs(rec).read_bytes()[:5] == b"%PDF-"          # a real PDF (reportlab, in-process)


def test_make_presentation_pptx():
    rec = d.make_presentation(_deck_spec())
    assert rec["ok"] is True and rec["kind"] == "presentation" and rec["path"].endswith(".pptx")
    from pptx import Presentation
    prs = Presentation(str(_abs(rec)))
    assert len(prs.slides) == 2


def test_make_workbook_xlsx_live_formula():
    rec = d.make_workbook(_workbook_spec())
    assert rec["ok"] is True and rec["path"].endswith(".xlsx")
    # the mass cell resolved from a tool result; the total is a REAL live formula.
    assert any(r["key"] == "mass_g" for r in rec["source_refs"])
    from openpyxl import load_workbook
    ws = load_workbook(str(_abs(rec))).active
    assert ws["B2"].value == 214.8
    assert isinstance(ws["D2"].value, str) and ws["D2"].value.startswith("=")


# ---------------------------------------------------------- (b) provenance violation

def test_literal_in_load_bearing_slot_is_error():
    spec = _report_spec("docx")
    spec["outline"][1]["slots"]["max_stress"] = 182.4     # a bare numeral, not a ref
    out = d.make_report(spec)
    assert out["ok"] is False and out["reason"] == "provenance_violation"
    # failed BEFORE any file is written: nothing landed in the registry.
    assert d.list_artifacts()["count"] == 0


def test_numeral_in_load_bearing_workbook_column_is_error():
    spec = _workbook_spec()
    spec["outline"][0]["rows"][1][1] = 214.8              # typed number where a ref belongs
    out = d.make_workbook(spec)
    assert out["ok"] is False and out["reason"] == "provenance_violation"


def test_label_string_in_value_column_is_allowed():
    # A label cell (header "Mass (g)") in a load-bearing column is free prose, not a violation.
    rec = d.make_workbook(_workbook_spec())
    assert rec["ok"] is True


# ---------------------------------------------------------- (c) slot_unresolved

def test_slot_unresolved_missing_result_id():
    spec = _report_spec("docx")
    spec["outline"][1]["slots"]["min_fos"]["from"]["result"] = "ghost"
    out = d.make_report(spec)
    assert out["ok"] is False and out["reason"] == "slot_unresolved"


def test_slot_unresolved_missing_key():
    spec = _report_spec("docx")
    spec["outline"][1]["slots"]["min_fos"]["from"]["key"] = "not_a_key"
    out = d.make_report(spec)
    assert out["ok"] is False and out["reason"] == "slot_unresolved"


def test_slot_unresolved_ambiguous_tool():
    spec = _report_spec("docx")
    # two results share the tool and the ref omits a result id -> ambiguous.
    spec["attached_results"]["fea2"] = dict(_ATTACHED["fea1"])
    spec["outline"][1]["slots"]["max_stress"]["from"].pop("result")
    out = d.make_report(spec)
    assert out["ok"] is False and out["reason"] == "slot_unresolved"


# ---------------------------------------------------------- invalid_spec / unsupported_format

def test_kind_mismatch_is_invalid_spec():
    out = d.make_report(_deck_spec())                     # a presentation spec into make_report
    assert out["ok"] is False and out["reason"] == "invalid_spec"


def test_unknown_template_is_invalid_spec():
    spec = _report_spec("docx")
    spec["template"] = "no_such_template"
    assert d.make_report(spec)["reason"] == "invalid_spec"


def test_unsupported_format():
    spec = _report_spec("pptx")                            # pptx not allowed for make_report
    assert d.make_report(spec)["reason"] == "unsupported_format"


# ---------------------------------------------------------- (d) render_document

def _fake_soffice(monkeypatch):
    """Mock LibreOffice: pretend it's on PATH and have subprocess.run emit a real 1-page PDF."""
    monkeypatch.setattr(engines, "soffice_bin", lambda: "/usr/bin/soffice-mock")

    def fake_run(cmd, *a, **k):
        outdir = Path(cmd[cmd.index("--outdir") + 1])
        src = Path(cmd[-1])
        pdf = outdir / (src.stem + ".pdf")
        from reportlab.lib.pagesizes import letter
        from reportlab.pdfgen import canvas
        c = canvas.Canvas(str(pdf), pagesize=letter)
        c.drawString(72, 720, f"preview of {src.name}")
        c.showPage()
        c.save()

        class _P:
            returncode = 0
            stderr = ""
        return _P()

    monkeypatch.setattr(engines.subprocess, "run", fake_run)


def test_render_office_with_mocked_converter(monkeypatch):
    rec = d.make_report(_report_spec("docx"))
    _fake_soffice(monkeypatch)
    out = d.render_document(rec["path"])
    assert out["ok"] is True and out["kind"] == "docx"
    assert out["page_count"] >= 1
    img = out["pages"][0]["image"]
    assert max(img["w"], img["h"]) <= engines.MAX_LONG_EDGE
    assert img["bytes"] < engines.MAX_BYTES
    assert out["resource_link"].startswith("file://")


def test_render_office_degrades_without_libreoffice(monkeypatch):
    rec = d.make_report(_report_spec("docx"))
    monkeypatch.setattr(engines, "soffice_bin", lambda: None)   # LibreOffice not on PATH
    out = d.render_document(rec["path"])
    assert out["ok"] is False and out["reason"] == "converter_missing"
    assert "libreoffice" in out["hint"].lower()
    # the artifact file itself is still on disk — only the preview degraded.
    assert _abs(rec).exists()


def test_render_pdf_native_no_subprocess(monkeypatch):
    rec = d.make_report(_report_spec("pdf"))
    # even with LibreOffice absent, the PDF path is native (pypdfium2) and must work.
    monkeypatch.setattr(engines, "soffice_bin", lambda: None)
    out = d.render_document(rec["path"])
    assert out["ok"] is True and out["kind"] == "pdf" and out["page_count"] >= 1
    img = out["pages"][0]["image"]
    assert max(img["w"], img["h"]) <= engines.MAX_LONG_EDGE and img["bytes"] < engines.MAX_BYTES


def test_render_csv_native(monkeypatch):
    monkeypatch.setattr(engines, "soffice_bin", lambda: None)
    csv = project.artifacts_dir() / "data.csv"
    csv.write_text("a,b,c\n1,2,3\n4,5,6\n", encoding="utf-8")
    out = d.render_document("artifacts/data.csv")
    assert out["ok"] is True and out["kind"] == "csv" and out["page_count"] == 1


def test_render_missing_file_and_unsupported():
    assert d.render_document("artifacts/nope.docx")["reason"] == "artifact_not_found"
    bad = project.artifacts_dir() / "thing.zzz"
    bad.write_text("x", encoding="utf-8")
    assert d.render_document("artifacts/thing.zzz")["reason"] == "unsupported_format"


# ---------------------------------------------------------- (e) registry round-trip + filter

def test_registry_roundtrip_and_run_id_filter():
    a = _report_spec("docx"); a["run_id"] = "runA"
    b = _report_spec("docx"); b["run_id"] = "runA"
    c = _deck_spec(); c["run_id"] = "runB"
    ra = d.make_report(a); d.make_report(b); rc = d.make_presentation(c)
    assert ra["ok"] and rc["ok"]

    allrec = d.list_artifacts()
    assert allrec["ok"] is True and allrec["count"] == 3
    # newest-first: the last-created artifact leads.
    assert allrec["artifacts"][0]["artifact_id"] == rc["artifact_id"]

    only_a = d.list_artifacts("runA")
    assert only_a["count"] == 2 and all(r["run_id"] == "runA" for r in only_a["artifacts"])

    # unknown run_id -> count 0, empty list, NOT an error.
    ghost = d.list_artifacts("nope")
    assert ghost["ok"] is True and ghost["count"] == 0 and ghost["artifacts"] == []


# ---------------------------------------------------------- registration + guide

def test_registration_and_guide():
    class _Stub:
        def __init__(self):
            self.names = []

        def tool(self):
            def deco(fn):
                self.names.append(fn.__name__)
                return fn
            return deco

    stub = _Stub()
    d.register_deliverables_tools(stub)
    assert set(stub.names) == {"make_presentation", "make_report", "make_workbook",
                               "make_work_instruction", "make_weldment_traveler",
                               "render_document", "list_artifacts"}
    assert "provenance" in d.DELIVERABLES_GUIDE.lower()


if __name__ == "__main__":
    raise SystemExit(pytest.main([__file__, "-q"]))
